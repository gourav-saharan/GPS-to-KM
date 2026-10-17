"""
GPS Travel History → Per-Day KM Report Generator with Vehicle Logbook
===================================================================
A single-page Python application that processes GPS Travel History Excel files
exported from vehicle tracking systems, and generates per-vehicle, per-day
KM summary reports with persistent odometer logbook management.

Author: Gourav Saharan
Version: 2.0 (Logbook Edition)
"""

import os
import sys
import json
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from datetime import datetime, timedelta
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DND_AVAILABLE = False
try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    DND_AVAILABLE = True
except ImportError:
    pass




class OdometerManager:
    """Manages persistent vehicle odometer readings stored in JSON."""

    STORAGE_FILE = "vehicle_odometers.json"

    def __init__(self, storage_dir=None):
        if storage_dir is None:
            storage_dir = os.path.dirname(os.path.abspath(__file__))
        self.file_path = os.path.join(storage_dir, self.STORAGE_FILE)
        self.data = self.load()

    def load(self):
        """Load odometer data from JSON file."""
        if os.path.exists(self.file_path):
            try:
                with open(self.file_path, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception:
                return {}
        return {}

    def save(self):
        """Save odometer data to JSON file."""
        try:
            with open(self.file_path, 'w', encoding='utf-8') as f:
                json.dump(self.data, f, indent=2)
            return True
        except Exception:
            return False

    def get_odometer(self, vehicle_name, default=None):
        """Get current saved odometer for a vehicle."""
        if vehicle_name in self.data:
            entry = self.data[vehicle_name]
            if isinstance(entry, dict):
                return float(entry.get('odometer', 0.0))
            elif isinstance(entry, (int, float)):
                return float(entry)
        return float(default) if default is not None else None

    def set_odometer(self, vehicle_name, odometer_val, last_file=""):
        """Set/update saved odometer for a vehicle."""
        try:
            val = float(odometer_val)
        except (ValueError, TypeError):
            val = 0.0

        if vehicle_name not in self.data or not isinstance(self.data[vehicle_name], dict):
            self.data[vehicle_name] = {}

        self.data[vehicle_name]['odometer'] = round(val, 2)
        self.data[vehicle_name]['last_updated'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        if last_file:
            self.data[vehicle_name]['last_file'] = last_file
        self.save()

    def delete_vehicle(self, vehicle_name):
        """Delete a vehicle record."""
        if vehicle_name in self.data:
            del self.data[vehicle_name]
            self.save()

    def get_all_vehicles(self):
        """Get all vehicle records."""
        return dict(self.data)

class TravelDataProcessor:
    """Processes GPS Travel History Excel files and computes per-day summaries."""

    EXPECTED_HEADERS = ['Vehicle', 'Status', 'Start Time', 'End Time', 'Distance']
    HEADER_ROW = 5
    DATA_START_ROW = 6

    DATE_FORMAT = '%d-%m-%Y %I:%M:%S %p'

    def __init__(self, filepath, progress_callback=None):
        self.filepath = filepath
        self.progress_callback = progress_callback
        self.vehicles = OrderedDict()  
        self.errors = []
        self.total_rows = 0

    def _update_progress(self, value, message=""):
        """Thread-safe progress update."""
        if self.progress_callback:
            self.progress_callback(value, message)

    def validate_file(self):
        """Validate that the file is a proper Travel History Excel file."""
        if not os.path.exists(self.filepath):
            return False, "File does not exist."
        if not self.filepath.lower().endswith('.xlsx'):
            return False, "File must be an .xlsx Excel file."
        try:
            import warnings
            warnings.filterwarnings("ignore", category=UserWarning)

            wb = openpyxl.load_workbook(self.filepath, read_only=True, data_only=True)
            ws = wb.active
            if ws is None:
                wb.close()
                return False, "Excel file has no sheets."

            header_values = []
            for row in ws.iter_rows(min_row=self.HEADER_ROW, max_row=self.HEADER_ROW,
                                     max_col=20, values_only=True):
                for val in row:
                    if val:
                        header_values.append(str(val).strip())

            wb.close()

            missing = []
            for expected in self.EXPECTED_HEADERS:
                if not any(expected.lower() in h.lower() for h in header_values):
                    missing.append(expected)

            if missing:
                return False, f"Missing expected columns: {', '.join(missing)}. This may not be a Travel History file."

            return True, "File is valid."
        except Exception as e:
            return False, f"Error reading file: {str(e)}"

    def parse_datetime(self, dt_string):
        """Parse datetime string in DD-MM-YYYY HH:MM:SS AM/PM format."""
        if dt_string is None:
            return None
        try:
            dt_str = str(dt_string).strip()
            return datetime.strptime(dt_str, self.DATE_FORMAT)
        except ValueError:
            alt_formats = [
                '%d-%m-%Y %H:%M:%S',
                '%Y-%m-%d %H:%M:%S',
                '%d/%m/%Y %I:%M:%S %p',
                '%d/%m/%Y %H:%M:%S',
            ]
            for fmt in alt_formats:
                try:
                    return datetime.strptime(dt_str, fmt)
                except ValueError:
                    continue
            return None

    def parse_duration_to_seconds(self, duration_str):
        """Parse HH:MM:SS duration string to total seconds."""
        if duration_str is None:
            return 0
        try:
            parts = str(duration_str).strip().split(':')
            if len(parts) == 3:
                h, m, s = int(parts[0]), int(parts[1]), int(parts[2])
                return h * 3600 + m * 60 + s
            return 0
        except (ValueError, IndexError):
            return 0

    def seconds_to_hms(self, total_seconds):
        """Convert total seconds to HH:MM:SS string."""
        if total_seconds <= 0:
            return "00:00:00"
        h = int(total_seconds // 3600)
        m = int((total_seconds % 3600) // 60)
        s = int(total_seconds % 60)
        return f"{h:02d}:{m:02d}:{s:02d}"

    def _safe_float(self, value):
        """Safely convert a value to float."""
        if value is None:
            return 0.0
        try:
            return float(value)
        except (ValueError, TypeError):
            return 0.0

    def read_and_process(self):
        """Read the Excel file and group data by vehicle."""
        self._update_progress(5, "Opening Excel file...")

        import warnings
        warnings.filterwarnings("ignore", category=UserWarning)

        wb = openpyxl.load_workbook(self.filepath, read_only=True, data_only=True)
        ws = wb.active

        self.total_rows = ws.max_row - self.DATA_START_ROW + 1
        self._update_progress(10, f"Reading {self.total_rows} data rows...")

        IDX_VEHICLE = 3       # D
        IDX_STATUS = 4        # E
        IDX_START_TIME = 5    # F
        IDX_START_LOC = 6     # G
        IDX_END_TIME = 7      # H
        IDX_END_LOC = 8       # I
        IDX_AVG_SPEED = 9     # J
        IDX_MAX_SPEED = 10    # K
        IDX_LAST_DIST = 11    # L
        IDX_DISTANCE = 12     # M
        IDX_DURATION = 13     # N

        current_vehicle = None
        rows_processed = 0

        for row_values in ws.iter_rows(min_row=self.DATA_START_ROW,
                                        max_col=20,
                                        values_only=True):
            rows_processed += 1

            if rows_processed % 1000 == 0:
                pct = 10 + int((rows_processed / max(self.total_rows, 1)) * 40)
                self._update_progress(min(pct, 49), f"Reading row {rows_processed}/{self.total_rows}...")

            vehicle = row_values[IDX_VEHICLE]
            status = row_values[IDX_STATUS]
            start_time_str = row_values[IDX_START_TIME]
            start_location = row_values[IDX_START_LOC]
            end_time_str = row_values[IDX_END_TIME]
            end_location = row_values[IDX_END_LOC]
            avg_speed = row_values[IDX_AVG_SPEED]
            max_speed = row_values[IDX_MAX_SPEED]
            last_distance = row_values[IDX_LAST_DIST]
            cumulative_dist = row_values[IDX_DISTANCE]
            duration_str = row_values[IDX_DURATION]

            if vehicle:
                current_vehicle = str(vehicle).strip()

            if current_vehicle is None:
                continue

            start_dt = self.parse_datetime(start_time_str)
            end_dt = self.parse_datetime(end_time_str)

            status_str = str(status).strip() if status else ""
            duration_secs = self.parse_duration_to_seconds(duration_str)

            row_data = {
                'vehicle': current_vehicle,
                'status': status_str,
                'start_dt': start_dt,
                'start_location': str(start_location).strip() if start_location else "",
                'end_dt': end_dt,
                'end_location': str(end_location).strip() if end_location else "",
                'avg_speed': self._safe_float(avg_speed),
                'max_speed': self._safe_float(max_speed),
                'last_distance': self._safe_float(last_distance),
                'cumulative_distance': self._safe_float(cumulative_dist),
                'duration_seconds': duration_secs,
            }

            if current_vehicle not in self.vehicles:
                self.vehicles[current_vehicle] = []
            self.vehicles[current_vehicle].append(row_data)

        wb.close()
        self._update_progress(50, f"Read complete. Found {len(self.vehicles)} vehicles.")
        return len(self.vehicles)

    def get_vehicle_summary_stats(self):
        """Returns summary stats per vehicle from parsed file."""
        stats = OrderedDict()
        for vehicle_name, rows in self.vehicles.items():
            if not rows:
                continue
            first_row = rows[0]
            last_row = rows[-1]
            start_gps_km = first_row['cumulative_distance'] - first_row['last_distance']
            end_gps_km = last_row['cumulative_distance']
            total_gps_km = max(0.0, end_gps_km - start_gps_km)
            stats[vehicle_name] = {
                'start_gps_km': round(start_gps_km, 2),
                'end_gps_km': round(end_gps_km, 2),
                'total_gps_km': round(total_gps_km, 2),
                'row_count': len(rows)
            }
        return stats

    def compute_daily_summary(self, vehicle_name, optional_columns, starting_odometer=None):
        """Compute per-day summary for a specific vehicle, applying odometer offset if provided."""
        rows = self.vehicles.get(vehicle_name, [])
        if not rows:
            return []

        # Determine odometer offset
        odometer_offset = 0.0
        if starting_odometer is not None:
            try:
                start_val = float(starting_odometer)
                first_row_file = rows[0]
                file_start_gps_km = first_row_file['cumulative_distance'] - first_row_file['last_distance']
                odometer_offset = start_val - file_start_gps_km
            except (ValueError, TypeError):
                odometer_offset = 0.0

        daily_groups = OrderedDict()
        for row in rows:
            if row['start_dt'] is None:
                continue
            date_key = row['start_dt'].strftime('%d-%m-%Y')
            if date_key not in daily_groups:
                daily_groups[date_key] = []
            daily_groups[date_key].append(row)

        sorted_dates = sorted(daily_groups.keys(),
                              key=lambda d: datetime.strptime(d, '%d-%m-%Y'))

        summaries = []
        for date_str in sorted_dates:
            day_rows = daily_groups[date_str]
            summary = self._compute_single_day(date_str, day_rows, optional_columns, odometer_offset)
            summaries.append(summary)

        return summaries

    def _compute_single_day(self, date_str, day_rows, optional_columns, odometer_offset=0.0):
        """Compute summary for a single day's data with odometer offset."""
        summary = {'date': date_str}

        first_row = day_rows[0]
        last_row = day_rows[-1]

        raw_start_km = first_row['cumulative_distance'] - first_row['last_distance']
        raw_end_km = last_row['cumulative_distance']

        actual_start_km = raw_start_km + odometer_offset
        actual_end_km = raw_end_km + odometer_offset

        summary['start_km'] = round(actual_start_km, 2)
        summary['end_km'] = round(actual_end_km, 2)
        summary['total_km'] = round(summary['end_km'] - summary['start_km'], 2)

        if 'running_time' in optional_columns:
            running_secs = sum(r['duration_seconds'] for r in day_rows
                               if r['status'].lower() == 'running')
            summary['running_time'] = self.seconds_to_hms(running_secs)

        if 'idle_time' in optional_columns:
            idle_secs = sum(r['duration_seconds'] for r in day_rows
                            if r['status'].lower() == 'idle')
            summary['idle_time'] = self.seconds_to_hms(idle_secs)

        if 'stop_time' in optional_columns:
            stop_secs = sum(r['duration_seconds'] for r in day_rows
                            if r['status'].lower() == 'stopped')
            summary['stop_time'] = self.seconds_to_hms(stop_secs)

        if 'max_speed' in optional_columns:
            max_spd = max((r['max_speed'] for r in day_rows), default=0.0)
            summary['max_speed'] = round(max_spd, 1)

        if 'avg_speed' in optional_columns:
            total_dist = sum(r['last_distance'] for r in day_rows if r['status'].lower() == 'running')
            if total_dist > 0:
                weighted_sum = sum(r['avg_speed'] * r['last_distance']
                                   for r in day_rows if r['status'].lower() == 'running')
                summary['avg_speed'] = round(weighted_sum / total_dist, 1)
            else:
                summary['avg_speed'] = 0.0

        if 'start_location' in optional_columns:
            summary['start_location'] = first_row['start_location']

        if 'end_location' in optional_columns:
            summary['end_location'] = last_row['end_location']

        if 'trip_count' in optional_columns:
            summary['trip_count'] = sum(1 for r in day_rows
                                         if r['status'].lower() == 'running')

        return summary


# =============================================================================
# EXCEL OUTPUT WRITER
# =============================================================================

class ReportWriter:
    """Writes per-vehicle daily summaries to a formatted Excel workbook."""

    HEADER_BG = "1B2A4A"       # Deep navy blue
    HEADER_FG = "FFFFFF"       # White
    SUBHEADER_BG = "2D4A7A"    # Medium blue
    ALT_ROW_BG1 = "F0F4FA"    # Very light blue
    ALT_ROW_BG2 = "FFFFFF"    # White
    TOTAL_ROW_BG = "E8F0FE"   # Light blue
    BORDER_COLOR = "B0C4DE"    # Light steel blue
    ACCENT_COLOR = "3B82F6"    # Bright blue accent

    DEFAULT_COLUMNS = [
        ('date', 'Date', 14),
        ('start_km', 'Start KM', 14),
        ('end_km', 'End KM', 14),
        ('total_km', 'Total KM', 12),
    ]

    OPTIONAL_COLUMN_DEFS = {
        'running_time': ('Running Time', 14),
        'idle_time': ('Idle Time', 14),
        'stop_time': ('Stop Time', 14),
        'max_speed': ('Max Speed (km/h)', 18),
        'avg_speed': ('Avg Speed (km/h)', 18),
        'start_location': ('Start Location', 35),
        'end_location': ('End Location', 35),
        'trip_count': ('Trip Count', 12),
    }

    def __init__(self, output_path, optional_columns):
        self.output_path = output_path
        self.optional_columns = optional_columns
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)

    def _get_column_list(self):
        """Get the full ordered list of columns to write."""
        cols = list(self.DEFAULT_COLUMNS)
        for key in self.optional_columns:
            if key in self.OPTIONAL_COLUMN_DEFS:
                name, width = self.OPTIONAL_COLUMN_DEFS[key]
                cols.append((key, name, width))
        return cols

    def _sanitize_sheet_name(self, name):
        """Sanitize vehicle name for use as Excel sheet name."""
        sanitized = name.replace('\\', '-').replace('/', '-').replace('*', '')\
                        .replace('?', '').replace(':', '-').replace('[', '(')\
                        .replace(']', ')')
        return sanitized[:31]

    def write_vehicle_sheet(self, vehicle_name, summaries):
        """Write a single vehicle's summary data to a new sheet."""
        sheet_name = self._sanitize_sheet_name(vehicle_name)
        ws = self.wb.create_sheet(title=sheet_name)

        columns = self._get_column_list()
        num_cols = len(columns)

        # --- Styles ---
        header_font = Font(name='Segoe UI', bold=True, size=11, color=self.HEADER_FG)
        header_fill = PatternFill(start_color=self.HEADER_BG, end_color=self.HEADER_BG, fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        title_font = Font(name='Segoe UI', bold=True, size=14, color=self.HEADER_BG)
        title_alignment = Alignment(horizontal='left', vertical='center')

        data_font = Font(name='Segoe UI', size=10)
        data_alignment_center = Alignment(horizontal='center', vertical='center')
        data_alignment_left = Alignment(horizontal='left', vertical='center')
        data_alignment_right = Alignment(horizontal='right', vertical='center')

        alt_fill1 = PatternFill(start_color=self.ALT_ROW_BG1, end_color=self.ALT_ROW_BG1, fill_type='solid')
        alt_fill2 = PatternFill(start_color=self.ALT_ROW_BG2, end_color=self.ALT_ROW_BG2, fill_type='solid')
        total_fill = PatternFill(start_color=self.TOTAL_ROW_BG, end_color=self.TOTAL_ROW_BG, fill_type='solid')
        total_font = Font(name='Segoe UI', bold=True, size=10, color=self.HEADER_BG)

        thin_border = Border(
            left=Side(style='thin', color=self.BORDER_COLOR),
            right=Side(style='thin', color=self.BORDER_COLOR),
            top=Side(style='thin', color=self.BORDER_COLOR),
            bottom=Side(style='thin', color=self.BORDER_COLOR),
        )

        # Title Row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        title_cell = ws.cell(row=1, column=1, value=f"📋 {vehicle_name}")
        title_cell.font = title_font
        title_cell.alignment = title_alignment
        ws.row_dimensions[1].height = 35

        # Summary info row
        if summaries:
            date_range = f"Period: {summaries[0]['date']} to {summaries[-1]['date']}  |  Days: {len(summaries)}"
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
            info_cell = ws.cell(row=2, column=1, value=date_range)
            info_cell.font = Font(name='Segoe UI', size=10, italic=True, color="6B7280")
            info_cell.alignment = Alignment(horizontal='left', vertical='center')
            ws.row_dimensions[2].height = 22

        # Header Row (row 4)
        header_row = 4
        ws.row_dimensions[header_row].height = 30
        for col_idx, (key, name, width) in enumerate(columns, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Data Rows
        data_start_row = header_row + 1
        total_km_sum = 0.0

        for row_offset, summary in enumerate(summaries):
            row_num = data_start_row + row_offset
            fill = alt_fill1 if row_offset % 2 == 0 else alt_fill2

            for col_idx, (key, name, width) in enumerate(columns, 1):
                value = summary.get(key, "")
                cell = ws.cell(row=row_num, column=col_idx, value=value)
                cell.font = data_font
                cell.fill = fill
                cell.border = thin_border

                if key == 'date':
                    cell.alignment = data_alignment_center
                elif key in ('start_location', 'end_location'):
                    cell.alignment = data_alignment_left
                elif key in ('start_km', 'end_km', 'total_km', 'max_speed', 'avg_speed', 'trip_count'):
                    cell.alignment = data_alignment_right
                    if isinstance(value, (int, float)):
                        cell.number_format = '#,##0.00' if key in ('start_km', 'end_km', 'total_km') else '#,##0.0'
                        if key == 'trip_count':
                            cell.number_format = '#,##0'
                else:
                    cell.alignment = data_alignment_center

            total_km_sum += summary.get('total_km', 0.0)

        # Total Row
        if summaries:
            total_row = data_start_row + len(summaries)
            ws.row_dimensions[total_row].height = 28

            total_label_cell = ws.cell(row=total_row, column=1, value="TOTAL")
            total_label_cell.font = total_font
            total_label_cell.fill = total_fill
            total_label_cell.alignment = data_alignment_center
            total_label_cell.border = thin_border

            for col_idx in range(2, num_cols + 1):
                cell = ws.cell(row=total_row, column=col_idx)
                cell.fill = total_fill
                cell.border = thin_border
                cell.font = total_font

            total_km_col = None
            for col_idx, (key, name, width) in enumerate(columns, 1):
                if key == 'total_km':
                    total_km_col = col_idx
                    break

            if total_km_col:
                total_km_cell = ws.cell(row=total_row, column=total_km_col,
                                         value=round(total_km_sum, 2))
                total_km_cell.font = total_font
                total_km_cell.fill = total_fill
                total_km_cell.alignment = Alignment(horizontal='right', vertical='center')
                total_km_cell.number_format = '#,##0.00'
                total_km_cell.border = thin_border

        ws.freeze_panes = ws.cell(row=data_start_row, column=1)

    def save(self):
        """Save the workbook to disk."""
        self.wb.save(self.output_path)
        self.wb.close()


# =============================================================================
# MODAL DIALOG: ODOMETER MANAGER
# =============================================================================

class OdometerManagerDialog(tk.Toplevel):
    """Modal dialog for managing saved vehicle odometer readings."""
    def __init__(self, parent, odometer_mgr, on_save_callback=None):
        super().__init__(parent)
        self.parent = parent
        self.odometer_mgr = odometer_mgr
        self.on_save_callback = on_save_callback

        self.title("⚙️ Manage Saved Vehicle Odometers")
        self.geometry("560x520")
        self.minsize(480, 420)
        self.configure(bg="#F8FAFC")
        self.transient(parent)
        self.grab_set()

        # Center on parent
        self.update_idletasks()
        x = parent.winfo_x() + (parent.winfo_width() // 2) - (560 // 2)
        y = parent.winfo_y() + (parent.winfo_height() // 2) - (520 // 2)
        self.geometry(f"+{max(0, x)}+{max(0, y)}")

        self.entries = {}
        self._build_ui()

    def _build_ui(self):
        hdr = tk.Frame(self, bg="#F8FAFC", padx=20, pady=16)
        hdr.pack(fill='x')

        title = tk.Label(
            hdr,
            text="📋 Saved Vehicle Odometer Logbook",
            font=("Segoe UI", 14, "bold"),
            fg="#1E293B",
            bg="#F8FAFC"
        )
        title.pack(anchor='w')

        desc = tk.Label(
            hdr,
            text="View and update starting odometer values stored in the application.",
            font=("Segoe UI", 9),
            fg="#64748B",
            bg="#F8FAFC"
        )
        desc.pack(anchor='w', pady=(2, 0))

        # Table container
        card = tk.Frame(self, bg="#CBD5E1", padx=1, pady=1)
        card.pack(fill='both', expand=True, padx=20, pady=(0, 12))

        inner_bg = tk.Frame(card, bg="#FFFFFF")
        inner_bg.pack(fill='both', expand=True)

        canvas = tk.Canvas(inner_bg, bg="#FFFFFF", highlightthickness=0)
        scrollbar = ttk.Scrollbar(inner_bg, orient="vertical", command=canvas.yview)
        scroll_frame = tk.Frame(canvas, bg="#FFFFFF")

        scroll_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        canvas.create_window((0, 0), window=scroll_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True, padx=8, pady=8)
        scrollbar.pack(side="right", fill="y")

        vehicles = self.odometer_mgr.get_all_vehicles()
        if not vehicles:
            no_data = tk.Label(
                scroll_frame,
                text="No saved vehicles yet.\nVehicles will automatically save here after report generation.",
                font=("Segoe UI", 10, "italic"),
                fg="#94A3B8",
                bg="#FFFFFF",
                pady=30
            )
            no_data.pack()
        else:
            th_frame = tk.Frame(scroll_frame, bg="#F1F5F9", pady=6, padx=8)
            th_frame.pack(fill='x', pady=(0, 8))

            tk.Label(th_frame, text="Vehicle Name", font=("Segoe UI", 9, "bold"), fg="#475569", bg="#F1F5F9", width=24, anchor='w').pack(side='left')
            tk.Label(th_frame, text="Current Odometer (KM)", font=("Segoe UI", 9, "bold"), fg="#475569", bg="#F1F5F9", width=20, anchor='w').pack(side='left', padx=(8, 0))
            tk.Label(th_frame, text="Action", font=("Segoe UI", 9, "bold"), fg="#475569", bg="#F1F5F9", anchor='e').pack(side='right')

            for v_name, v_info in vehicles.items():
                row = tk.Frame(scroll_frame, bg="#FFFFFF", pady=4, padx=8)
                row.pack(fill='x')

                v_label = tk.Label(
                    row,
                    text=v_name,
                    font=("Segoe UI", 9, "bold"),
                    fg="#1E293B",
                    bg="#FFFFFF",
                    width=24,
                    anchor='w'
                )
                v_label.pack(side='left')

                val = v_info.get('odometer', 0.0) if isinstance(v_info, dict) else v_info
                val_var = tk.StringVar(value=f"{float(val):.2f}")
                val_entry = tk.Entry(
                    row,
                    textvariable=val_var,
                    font=("Segoe UI", 9),
                    width=16,
                    bd=1,
                    relief="solid"
                )
                val_entry.pack(side='left', padx=(8, 0))
                self.entries[v_name] = val_var

                del_btn = tk.Button(
                    row,
                    text="🗑️ Delete",
                    font=("Segoe UI", 8),
                    fg="#DC2626",
                    bg="#FEE2E2",
                    bd=0,
                    cursor="hand2",
                    command=lambda name=v_name: self._delete_vehicle(name)
                )
                del_btn.pack(side='right')

        # Add New Vehicle Frame
        add_frame = tk.Frame(self, bg="#F8FAFC", padx=20)
        add_frame.pack(fill='x', pady=(0, 10))

        add_label = tk.Label(add_frame, text="+ Manually Add / Register Vehicle", font=("Segoe UI", 9, "bold"), fg="#2563EB", bg="#F8FAFC")
        add_label.pack(anchor='w', pady=(0, 4))

        add_inputs = tk.Frame(add_frame, bg="#F8FAFC")
        add_inputs.pack(fill='x')

        self.new_name_var = tk.StringVar()
        self.new_val_var = tk.StringVar()

        tk.Entry(add_inputs, textvariable=self.new_name_var, font=("Segoe UI", 9), width=24).pack(side='left')
        tk.Entry(add_inputs, textvariable=self.new_val_var, font=("Segoe UI", 9), width=14).pack(side='left', padx=6)

        add_btn = tk.Button(
            add_inputs,
            text="Add Vehicle",
            font=("Segoe UI", 9, "bold"),
            bg="#2563EB",
            fg="white",
            bd=0,
            padx=10,
            pady=2,
            cursor="hand2",
            command=self._add_vehicle
        )
        add_btn.pack(side='left')

        # Bottom Bar
        bbar = tk.Frame(self, bg="#F8FAFC", padx=20, pady=12)
        bbar.pack(fill='x', side='bottom')

        save_btn = tk.Button(
            bbar,
            text="💾 Save Changes",
            font=("Segoe UI", 10, "bold"),
            bg="#16A34A",
            fg="white",
            bd=0,
            padx=16,
            pady=6,
            cursor="hand2",
            command=self._save_changes
        )
        save_btn.pack(side='right')

        cancel_btn = tk.Button(
            bbar,
            text="Cancel",
            font=("Segoe UI", 10),
            bg="#E2E8F0",
            fg="#475569",
            bd=0,
            padx=14,
            pady=6,
            cursor="hand2",
            command=self.destroy
        )
        cancel_btn.pack(side='right', padx=(0, 8))

    def _delete_vehicle(self, vehicle_name):
        if messagebox.askyesno("Delete Vehicle", f"Remove saved odometer for '{vehicle_name}'?", parent=self):
            self.odometer_mgr.delete_vehicle(vehicle_name)
            if vehicle_name in self.entries:
                del self.entries[vehicle_name]
            self.destroy()
            OdometerManagerDialog(self.parent, self.odometer_mgr, self.on_save_callback)

    def _add_vehicle(self):
        name = self.new_name_var.get().strip()
        val_str = self.new_val_var.get().strip()
        if not name:
            messagebox.showwarning("Input Error", "Vehicle name cannot be empty.", parent=self)
            return
        try:
            val = float(val_str)
        except ValueError:
            messagebox.showwarning("Input Error", "Please enter a valid numeric odometer value.", parent=self)
            return

        self.odometer_mgr.set_odometer(name, val)
        self.destroy()
        OdometerManagerDialog(self.parent, self.odometer_mgr, self.on_save_callback)

    def _save_changes(self):
        for name, var in self.entries.items():
            try:
                val = float(var.get().strip())
                self.odometer_mgr.set_odometer(name, val)
            except ValueError:
                pass
        if self.on_save_callback:
            self.on_save_callback()
        self.destroy()


# =============================================================================
# TKINTER UI APPLICATION
# =============================================================================

class GPStoKMApp:
    """Main application UI."""

    BG_COLOR = "#F8FAFC"         # Light background
    CARD_BG = "#FFFFFF"          # White card background
    CARD_BORDER = "#E2E8F0"      # Soft gray border
    ACCENT = "#2563EB"           # Blue accent
    ACCENT_HOVER = "#1D4ED8"     # Darker blue on hover
    ACCENT_LIGHT = "#3B82F6"     # Lighter blue
    SUCCESS = "#16A34A"          # Green
    ERROR = "#DC2626"            # Red
    TEXT_PRIMARY = "#1E293B"     # Dark text
    TEXT_SECONDARY = "#64748B"   # Medium gray text
    TEXT_MUTED = "#94A3B8"       # Muted text
    DROPZONE_BG = "#F1F5F9"      # Drop zone light background
    DROPZONE_ACTIVE = "#DBEAFE"  # Drop zone active

    CHECKBOX_ITEMS = [
        ('running_time', 'Running Time', True),
        ('idle_time', 'Idle Time', True),
        ('stop_time', 'Stop Time', False),
        ('max_speed', 'Max Speed', False),
        ('avg_speed', 'Avg Speed', False),
        ('start_location', 'Start Location', False),
        ('end_location', 'End Location', False),
        ('trip_count', 'Trip Count', False),
    ]

    def __init__(self):
        if sys.platform == 'win32':
            try:
                import ctypes
                u32 = ctypes.windll.user32
                DESKTOP_ALL_ACCESS = 0x01FF
                hdesk = u32.OpenDesktopW('Default', 0, False, DESKTOP_ALL_ACCESS)
                if hdesk:
                    u32.SetThreadDesktop(hdesk)
            except Exception:
                pass

        if DND_AVAILABLE:
            self.root = TkinterDnD.Tk()
        else:
            self.root = tk.Tk()

        self.root.title("GPS to KM Report Generator & Odometer Logbook")
        self.root.geometry("700x820")
        self.root.minsize(620, 750)
        self.root.configure(bg=self.BG_COLOR)

        # Center window
        self.root.update_idletasks()
        x = (self.root.winfo_screenwidth() // 2) - (700 // 2)
        y = (self.root.winfo_screenheight() // 2) - (820 // 2)
        self.root.geometry(f"+{max(0, x)}+{max(0, y)}")

        # Bring window to front
        try:
            self.root.lift()
            self.root.attributes('-topmost', True)
            self.root.after_idle(self.root.attributes, '-topmost', False)
            self.root.after_idle(self.root.focus_force)
        except Exception:
            pass

        # Application Managers & Variables
        self.odometer_mgr = OdometerManager()
        self.selected_file = tk.StringVar(value="")
        self.status_text = tk.StringVar(value="Ready — Drop or select an Excel file to begin")
        self.progress_value = tk.DoubleVar(value=0)
        self.is_processing = False
        self.checkbox_vars = {}
        self.auto_update_var = tk.BooleanVar(value=True)

        # Detected Vehicles state: {v_name: {'start_var': StringVar, 'end_label': Label, 'total_km': float}}
        self.detected_vehicles_ui = {}
        self.parsed_vehicle_stats = {}

        # Build UI
        self._build_ui()

    def _build_ui(self):
        """Build the main UI layout."""
        main_frame = tk.Frame(self.root, bg=self.BG_COLOR)
        main_frame.pack(fill='both', expand=True, padx=24, pady=16)

        # --- Header ---
        self._build_header(main_frame)

        # --- Drop Zone ---
        self._build_dropzone(main_frame)

        # --- File Info ---
        self._build_file_info(main_frame)

        # --- Odometer Configuration Card (Dynamically updated) ---
        self._build_odometer_card(main_frame)

        # --- Optional Columns ---
        self._build_checkboxes(main_frame)

        # --- Generate Button ---
        self._build_generate_button(main_frame)

        # --- Progress & Status ---
        self._build_status(main_frame)

    def _build_header(self, parent):
        """Build the header with title and Odometer Manager button."""
        header_frame = tk.Frame(parent, bg=self.BG_COLOR)
        header_frame.pack(fill='x', pady=(0, 14))

        title_label = tk.Label(
            header_frame,
            text="🚗  GPS to KM Logbook Generator",
            font=("Segoe UI", 16, "bold"),
            fg=self.TEXT_PRIMARY,
            bg=self.BG_COLOR,
            anchor='w'
        )
        title_label.pack(side='left')

        mgr_btn = tk.Button(
            header_frame,
            text="⚙️ Manage Odometers",
            font=("Segoe UI", 9, "bold"),
            fg=self.ACCENT,
            bg="#EFF6FF",
            activebackground="#DBEAFE",
            activeforeground=self.ACCENT_HOVER,
            bd=1,
            relief="solid",
            padx=10,
            pady=4,
            cursor="hand2",
            command=self._open_odometer_dialog
        )
        mgr_btn.pack(side='right')

    def _build_dropzone(self, parent):
        """Build the drag-and-drop file zone."""
        card = tk.Frame(parent, bg=self.CARD_BORDER, padx=1, pady=1)
        card.pack(fill='x', pady=(0, 10))

        self.dropzone = tk.Frame(card, bg=self.DROPZONE_BG, cursor="hand2")
        self.dropzone.pack(fill='x')

        inner = tk.Frame(self.dropzone, bg=self.DROPZONE_BG)
        inner.pack(fill='x', padx=20, pady=20)

        icon_label = tk.Label(inner, text="📁", font=("Segoe UI", 28), bg=self.DROPZONE_BG)
        icon_label.pack()

        drop_text = tk.Label(
            inner,
            text="Drag & Drop Excel File Here" if DND_AVAILABLE else "Click to Select Excel File",
            font=("Segoe UI", 12, "bold"),
            fg=self.TEXT_PRIMARY,
            bg=self.DROPZONE_BG
        )
        drop_text.pack(pady=(2, 2))

        sub_text = tk.Label(
            inner,
            text="or click to browse (.xlsx Travel History file)",
            font=("Segoe UI", 9),
            fg=self.TEXT_SECONDARY,
            bg=self.DROPZONE_BG
        )
        sub_text.pack()

        for widget in [self.dropzone, inner, icon_label, drop_text, sub_text]:
            widget.bind("<Button-1>", self._browse_file)

        if DND_AVAILABLE:
            self.dropzone.drop_target_register(DND_FILES)
            self.dropzone.dnd_bind('<<DropEnter>>', self._on_drop_enter)
            self.dropzone.dnd_bind('<<DropLeave>>', self._on_drop_leave)
            self.dropzone.dnd_bind('<<Drop>>', self._on_drop)

    def _build_file_info(self, parent):
        """Build the file info display."""
        self.file_info_frame = tk.Frame(parent, bg=self.BG_COLOR)
        self.file_info_frame.pack(fill='x', pady=(0, 8))

        self.file_label = tk.Label(
            self.file_info_frame,
            text="",
            font=("Segoe UI", 10, "bold"),
            fg=self.ACCENT_LIGHT,
            bg=self.BG_COLOR,
            anchor='w'
        )
        self.file_label.pack(side='left', fill='x', expand=True)

        self.clear_btn = tk.Label(
            self.file_info_frame,
            text="✕ Clear",
            font=("Segoe UI", 9),
            fg=self.ERROR,
            bg=self.BG_COLOR,
            cursor="hand2"
        )
        self.clear_btn.bind("<Button-1>", self._clear_file)

    def _build_odometer_card(self, parent):
        """Build the interactive Detected Vehicles Odometer configuration card."""
        self.odo_card_frame = tk.Frame(parent, bg=self.BG_COLOR)
        self.odo_card_frame.pack(fill='x', pady=(0, 10))

        # We will dynamically populate this container when a file is loaded
        self.odo_inner_card = None

    def _update_odometer_card(self):
        """Refreshes the Detected Vehicles Odometer Configuration Card."""
        if self.odo_inner_card:
            self.odo_inner_card.destroy()

        if not self.parsed_vehicle_stats:
            return

        card_outer = tk.Frame(self.odo_card_frame, bg=self.CARD_BORDER, padx=1, pady=1)
        card_outer.pack(fill='x')
        self.odo_inner_card = card_outer

        card_bg = tk.Frame(card_outer, bg=self.CARD_BG, padx=14, pady=12)
        card_bg.pack(fill='x')

        # Section Header
        hdr_frame = tk.Frame(card_bg, bg=self.CARD_BG)
        hdr_frame.pack(fill='x', pady=(0, 8))

        tk.Label(
            hdr_frame,
            text="📋  Vehicle Odometer Configuration & Logbook",
            font=("Segoe UI", 11, "bold"),
            fg=self.TEXT_PRIMARY,
            bg=self.CARD_BG
        ).pack(side='left')

        tk.Label(
            hdr_frame,
            text="Start/End KM in report will be based on starting odometer values",
            font=("Segoe UI", 8, "italic"),
            fg=self.TEXT_MUTED,
            bg=self.CARD_BG
        ).pack(side='right')

        # Vehicle Rows Container
        rows_container = tk.Frame(card_bg, bg=self.CARD_BG)
        rows_container.pack(fill='x')

        # Table Subheaders
        sub_hdr = tk.Frame(rows_container, bg="#F1F5F9", pady=4, padx=8)
        sub_hdr.pack(fill='x', pady=(0, 6))

        tk.Label(sub_hdr, text="Vehicle Name", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#F1F5F9", width=24, anchor='w').pack(side='left')
        tk.Label(sub_hdr, text="Starting Odometer (KM)", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#F1F5F9", width=20, anchor='w').pack(side='left', padx=(6, 0))
        tk.Label(sub_hdr, text="File Travel", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#F1F5F9", width=12, anchor='w').pack(side='left')
        tk.Label(sub_hdr, text="Report End Odometer", font=("Segoe UI", 8, "bold"), fg="#475569", bg="#F1F5F9", anchor='e').pack(side='right')

        self.detected_vehicles_ui = {}

        for v_name, stats in self.parsed_vehicle_stats.items():
            row_frame = tk.Frame(rows_container, bg=self.CARD_BG, pady=3, padx=8)
            row_frame.pack(fill='x')

            v_label = tk.Label(
                row_frame,
                text=v_name,
                font=("Segoe UI", 9, "bold"),
                fg=self.TEXT_PRIMARY,
                bg=self.CARD_BG,
                width=24,
                anchor='w'
            )
            v_label.pack(side='left')

            # Default starting odometer: saved value, or file_start_gps_km if none saved
            saved_val = self.odometer_mgr.get_odometer(v_name, default=stats['start_gps_km'])
            start_var = tk.StringVar(value=f"{saved_val:.2f}")

            entry_box = tk.Entry(
                row_frame,
                textvariable=start_var,
                font=("Segoe UI", 9),
                width=16,
                bd=1,
                relief="solid"
            )
            entry_box.pack(side='left', padx=(6, 0))

            dist_label = tk.Label(
                row_frame,
                text=f"+ {stats['total_gps_km']:,.2f} KM",
                font=("Segoe UI", 8),
                fg=self.TEXT_SECONDARY,
                bg=self.CARD_BG,
                width=14,
                anchor='w'
            )
            dist_label.pack(side='left', padx=(8, 0))

            end_label = tk.Label(
                row_frame,
                text=f"= {saved_val + stats['total_gps_km']:,.2f} KM",
                font=("Segoe UI", 9, "bold"),
                fg=self.SUCCESS,
                bg=self.CARD_BG,
                anchor='e'
            )
            end_label.pack(side='right')

            # Live update trace
            def _make_trace_cb(var, end_lbl, total_km):
                def _cb(*args):
                    try:
                        s_val = float(var.get().strip())
                        end_lbl.config(text=f"= {s_val + total_km:,.2f} KM", fg=self.SUCCESS)
                    except ValueError:
                        end_lbl.config(text="Invalid Value", fg=self.ERROR)
                return _cb

            start_var.trace_add("write", _make_trace_cb(start_var, end_label, stats['total_gps_km']))

            self.detected_vehicles_ui[v_name] = {
                'start_var': start_var,
                'end_label': end_label,
                'total_km': stats['total_gps_km']
            }

        # Auto-update checkbox
        cb_frame = tk.Frame(card_bg, bg=self.CARD_BG)
        cb_frame.pack(fill='x', pady=(8, 0))

        cb_auto = tk.Checkbutton(
            cb_frame,
            text="🔄  Automatically update saved odometers after generating report",
            variable=self.auto_update_var,
            font=("Segoe UI", 9, "bold"),
            fg=self.ACCENT,
            bg=self.CARD_BG,
            activebackground=self.CARD_BG,
            activeforeground=self.ACCENT,
            selectcolor=self.CARD_BG,
            bd=0,
            cursor="hand2"
        )
        cb_auto.pack(anchor='w')

    def _build_checkboxes(self, parent):
        """Build optional columns checkbox panel."""
        section_frame = tk.Frame(parent, bg=self.BG_COLOR)
        section_frame.pack(fill='x', pady=(0, 4))

        section_label = tk.Label(
            section_frame,
            text="Optional Columns",
            font=("Segoe UI", 11, "bold"),
            fg=self.TEXT_PRIMARY,
            bg=self.BG_COLOR,
            anchor='w'
        )
        section_label.pack(side='left')

        hint_label = tk.Label(
            section_frame,
            text="(included in per-day vehicle sheets)",
            font=("Segoe UI", 9),
            fg=self.TEXT_MUTED,
            bg=self.BG_COLOR,
            anchor='w'
        )
        hint_label.pack(side='left', padx=(8, 0), pady=(2, 0))

        cb_card = tk.Frame(parent, bg=self.CARD_BORDER, padx=1, pady=1)
        cb_card.pack(fill='x', pady=(4, 14))

        cb_inner = tk.Frame(cb_card, bg=self.CARD_BG)
        cb_inner.pack(fill='x')

        grid_frame = tk.Frame(cb_inner, bg=self.CARD_BG)
        grid_frame.pack(fill='x', padx=16, pady=10)

        for idx, (key, label, default) in enumerate(self.CHECKBOX_ITEMS):
            var = tk.BooleanVar(value=default)
            self.checkbox_vars[key] = var

            row = idx // 3
            col = idx % 3

            cb_frame = tk.Frame(grid_frame, bg=self.CARD_BG)
            cb_frame.grid(row=row, column=col, sticky='w', padx=(0, 16), pady=3)

            cb = tk.Checkbutton(
                cb_frame,
                text=label,
                variable=var,
                font=("Segoe UI", 9),
                fg=self.TEXT_PRIMARY,
                bg=self.CARD_BG,
                selectcolor=self.CARD_BG,
                activebackground=self.CARD_BG,
                activeforeground=self.TEXT_PRIMARY,
                highlightthickness=0,
                bd=0,
                cursor="hand2"
            )
            cb.pack(anchor='w')

    def _build_generate_button(self, parent):
        """Build the generate report button."""
        btn_frame = tk.Frame(parent, bg=self.BG_COLOR)
        btn_frame.pack(fill='x', pady=(0, 14))

        self.generate_btn = tk.Canvas(
            btn_frame,
            height=46,
            bg=self.BG_COLOR,
            highlightthickness=0,
            cursor="hand2"
        )
        self.generate_btn.pack(fill='x')

        self.generate_btn.bind("<Configure>", self._draw_button)
        self.generate_btn.bind("<Button-1>", self._on_generate_click)
        self.generate_btn.bind("<Enter>", self._on_btn_enter)
        self.generate_btn.bind("<Leave>", self._on_btn_leave)
        self._btn_hover = False

    def _draw_button(self, event=None):
        """Draw the generate button with rounded corners."""
        canvas = self.generate_btn
        canvas.delete("all")
        w = canvas.winfo_width()
        h = canvas.winfo_height()
        color = self.ACCENT_HOVER if self._btn_hover else self.ACCENT

        if self.is_processing:
            color = self.TEXT_MUTED

        r = 8
        canvas.create_arc(0, 0, 2*r, 2*r, start=90, extent=90, fill=color, outline=color)
        canvas.create_arc(w-2*r, 0, w, 2*r, start=0, extent=90, fill=color, outline=color)
        canvas.create_arc(0, h-2*r, 2*r, h, start=180, extent=90, fill=color, outline=color)
        canvas.create_arc(w-2*r, h-2*r, w, h, start=270, extent=90, fill=color, outline=color)
        canvas.create_rectangle(r, 0, w-r, h, fill=color, outline=color)
        canvas.create_rectangle(0, r, w, h-r, fill=color, outline=color)

        text = "⏳ Processing Excel & Updating Logbook..." if self.is_processing else "🚀  Generate Logbook Excel Report"
        canvas.create_text(
            w // 2, h // 2,
            text=text,
            font=("Segoe UI", 12, "bold"),
            fill="white"
        )

    def _on_btn_enter(self, event):
        self._btn_hover = True
        self._draw_button()

    def _on_btn_leave(self, event):
        self._btn_hover = False
        self._draw_button()

    def _build_status(self, parent):
        """Build status bar and progress bar."""
        status_frame = tk.Frame(parent, bg=self.BG_COLOR)
        status_frame.pack(fill='x', side='bottom')

        style = ttk.Style()
        style.configure("Custom.Horizontal.TProgressbar",
                        troughcolor=self.CARD_BG,
                        background=self.ACCENT,
                        thickness=6)

        self.progress_bar = ttk.Progressbar(
            status_frame,
            variable=self.progress_value,
            maximum=100,
            style="Custom.Horizontal.TProgressbar",
            mode='determinate'
        )
        self.progress_bar.pack(fill='x', pady=(0, 6))

        self.status_label = tk.Label(
            status_frame,
            textvariable=self.status_text,
            font=("Segoe UI", 9),
            fg=self.TEXT_SECONDARY,
            bg=self.BG_COLOR,
            anchor='w'
        )
        self.status_label.pack(fill='x')

    # -------------------------------------------------------------------------
    # EVENT HANDLERS & FILE PARSING
    # -------------------------------------------------------------------------

    def _open_odometer_dialog(self):
        """Open the Manage Saved Odometers dialog."""
        OdometerManagerDialog(self.root, self.odometer_mgr, on_save_callback=self._update_odometer_card)

    def _browse_file(self, event=None):
        if self.is_processing:
            return

        filepath = filedialog.askopenfilename(
            title="Select Travel History Excel File",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            initialdir=os.path.expanduser("~\\Desktop")
        )
        if filepath:
            self._set_file(filepath)

    def _on_drop_enter(self, event):
        self.dropzone.configure(bg=self.DROPZONE_ACTIVE)
        for child in self.dropzone.winfo_children():
            self._set_bg_recursive(child, self.DROPZONE_ACTIVE)

    def _on_drop_leave(self, event):
        self.dropzone.configure(bg=self.DROPZONE_BG)
        for child in self.dropzone.winfo_children():
            self._set_bg_recursive(child, self.DROPZONE_BG)

    def _on_drop(self, event):
        self._on_drop_leave(event)
        filepath = event.data.strip()
        if filepath.startswith('{') and filepath.endswith('}'):
            filepath = filepath[1:-1]
        if filepath.startswith('"') and filepath.endswith('"'):
            filepath = filepath[1:-1]
        self._set_file(filepath)

    def _set_bg_recursive(self, widget, bg):
        try:
            widget.configure(bg=bg)
        except tk.TclError:
            pass
        for child in widget.winfo_children():
            self._set_bg_recursive(child, bg)

    def _set_file(self, filepath):
        """Set file and run quick scan to detect vehicles."""
        self.selected_file.set(filepath)
        self._update_file_display()

        filename = os.path.basename(filepath)
        if not filepath.lower().endswith('.xlsx'):
            self.status_text.set(f"⚠️  Invalid file type. Please select an .xlsx file.")
            self.status_label.configure(fg=self.ERROR)
            self.parsed_vehicle_stats = {}
            self._update_odometer_card()
            return

        self.status_text.set(f"⏳ Scanning vehicles in {filename}...")
        self.status_label.configure(fg=self.ACCENT)

        # Scan vehicles in background thread to keep UI fast
        threading.Thread(target=self._scan_file_vehicles, args=(filepath,), daemon=True).start()

    def _scan_file_vehicles(self, filepath):
        """Perform fast scan of vehicle names & distances."""
        try:
            processor = TravelDataProcessor(filepath)
            valid, msg = processor.validate_file()
            if not valid:
                def _err():
                    self.status_text.set(f"⚠️ {msg}")
                    self.status_label.configure(fg=self.ERROR)
                    self.parsed_vehicle_stats = {}
                    self._update_odometer_card()
                self.root.after(0, _err)
                return

            processor.read_and_process()
            stats = processor.get_vehicle_summary_stats()

            def _done():
                self.parsed_vehicle_stats = stats
                filename = os.path.basename(filepath)
                self.status_text.set(f"✅ Loaded {filename} ({len(stats)} vehicles detected)")
                self.status_label.configure(fg=self.SUCCESS)
                self._update_odometer_card()

            self.root.after(0, _done)
        except Exception as e:
            def _fail():
                self.status_text.set(f"⚠️ Error scanning file: {str(e)}")
                self.status_label.configure(fg=self.ERROR)
            self.root.after(0, _fail)

    def _clear_file(self, event=None):
        self.selected_file.set("")
        self.parsed_vehicle_stats = {}
        self.detected_vehicles_ui = {}
        self._update_odometer_card()
        self._update_file_display()
        self.progress_value.set(0)
        self.status_text.set("Ready — Drop or select an Excel file to begin")
        self.status_label.configure(fg=self.TEXT_SECONDARY)

    def _update_file_display(self):
        filepath = self.selected_file.get()
        if filepath:
            filename = os.path.basename(filepath)
            size_mb = os.path.getsize(filepath) / (1024 * 1024)
            self.file_label.configure(
                text=f"📄  {filename}  ({size_mb:.1f} MB)",
                fg=self.ACCENT_LIGHT
            )
            self.clear_btn.pack(side='right')
        else:
            self.file_label.configure(text="", fg=self.TEXT_SECONDARY)
            self.clear_btn.pack_forget()

    def _on_generate_click(self, event=None):
        """Handle Generate Report button click."""
        if self.is_processing:
            return

        filepath = self.selected_file.get()
        if not filepath:
            self.status_text.set("⚠️  Please select an Excel file first.")
            self.status_label.configure(fg=self.ERROR)
            return

        if not filepath.lower().endswith('.xlsx'):
            self.status_text.set("⚠️  Invalid file type. Please select an .xlsx file.")
            self.status_label.configure(fg=self.ERROR)
            return

        # Extract specified starting odometers from UI
        starting_odometers = {}
        for v_name, ui_info in self.detected_vehicles_ui.items():
            val_str = ui_info['start_var'].get().strip()
            try:
                starting_odometers[v_name] = float(val_str)
            except ValueError:
                self.status_text.set(f"⚠️ Invalid starting odometer for '{v_name}'")
                self.status_label.configure(fg=self.ERROR)
                return

        optional_cols = [key for key, var in self.checkbox_vars.items() if var.get()]
        do_auto_update = self.auto_update_var.get()

        default_name = os.path.splitext(os.path.basename(filepath))[0] + "_Logbook_Report.xlsx"
        output_path = filedialog.asksaveasfilename(
            title="Save Logbook Report As",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=default_name,
            initialdir=os.path.dirname(filepath)
        )

        if not output_path:
            return

        self.is_processing = True
        self._draw_button()

        thread = threading.Thread(
            target=self._process_file,
            args=(filepath, output_path, optional_cols, starting_odometers, do_auto_update),
            daemon=True
        )
        thread.start()

    def _update_progress_ui(self, value, message):
        def _update():
            self.progress_value.set(value)
            self.status_text.set(message)
            self.status_label.configure(fg=self.TEXT_SECONDARY)
        self.root.after(0, _update)

    def _process_file(self, filepath, output_path, optional_cols, starting_odometers, do_auto_update):
        """Process file and write output Excel report with odometer logic."""
        try:
            self._update_progress_ui(2, "Validating file format...")
            processor = TravelDataProcessor(filepath, self._update_progress_ui)

            valid, msg = processor.validate_file()
            if not valid:
                self._on_processing_error(f"Validation failed: {msg}")
                return

            num_vehicles = processor.read_and_process()
            if num_vehicles == 0:
                self._on_processing_error("No vehicle data found in the file.")
                return

            self._update_progress_ui(55, "Creating formatted Excel report sheets...")
            writer = ReportWriter(output_path, optional_cols)

            vehicle_names = list(processor.vehicles.keys())
            updated_odometer_summary = {}

            filename = os.path.basename(filepath)

            for v_idx, vehicle_name in enumerate(vehicle_names):
                pct = 55 + int((v_idx / len(vehicle_names)) * 35)
                self._update_progress_ui(pct, f"Processing: {vehicle_name}...")

                start_odo = starting_odometers.get(vehicle_name, None)
                summaries = processor.compute_daily_summary(vehicle_name, optional_cols, starting_odometer=start_odo)
                writer.write_vehicle_sheet(vehicle_name, summaries)

                if summaries:
                    final_end_km = summaries[-1]['end_km']
                    updated_odometer_summary[vehicle_name] = final_end_km

                    if do_auto_update:
                        self.odometer_mgr.set_odometer(vehicle_name, final_end_km, last_file=filename)

            self._update_progress_ui(95, "Saving Excel workbook...")
            writer.save()

            self._on_processing_complete(output_path, num_vehicles, updated_odometer_summary, do_auto_update)

        except Exception as e:
            self._on_processing_error(f"Error: {str(e)}")

    def _on_processing_complete(self, output_path, num_vehicles, updated_odometers, did_auto_update):
        """Called when report generation completes successfully."""
        def _update():
            self.is_processing = False
            self.progress_value.set(100)
            self.status_text.set(f"✅ Report saved & saved odometers updated ({num_vehicles} vehicles exported).")
            self.status_label.configure(fg=self.SUCCESS)
            self._draw_button()

            # Refresh odometer display if updated
            if did_auto_update:
                self.parsed_vehicle_stats = {}
                self._scan_file_vehicles(self.selected_file.get())

            odo_msg = ""
            if did_auto_update and updated_odometers:
                odo_msg = "\n\n🔄 Updated Saved Odometers for Future Use:\n"
                for v_name, new_val in updated_odometers.items():
                    odo_msg += f" • {v_name}: {new_val:,.2f} KM\n"

            result = messagebox.askyesno(
                "Logbook Report Generated Successfully! 🎉",
                f"Report saved to:\n{output_path}\n\n"
                f"Vehicles exported: {num_vehicles}"
                f"{odo_msg}\n"
                f"Would you like to open the output Excel file now?",
                parent=self.root
            )
            if result:
                os.startfile(output_path)

        self.root.after(0, _update)

    def _on_processing_error(self, error_msg):
        def _update():
            self.is_processing = False
            self.progress_value.set(0)
            self.status_text.set(f"❌  {error_msg}")
            self.status_label.configure(fg=self.ERROR)
            self._draw_button()
        self.root.after(0, _update)

    def run(self):
        """Start the application mainloop."""
        self.root.mainloop()


# =============================================================================
# APPLICATION ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    app = GPStoKMApp()
    app.run()
